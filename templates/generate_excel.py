#!/usr/bin/env python3
"""伏線管理テーブルの空テンプレートExcelを生成する"""

import sys
import os

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("エラー: openpyxl が必要です。 pip install openpyxl で導入してください。")
    sys.exit(1)


def generate(output_path, title="新連載"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "伏線管理テーブル"

    # Styles
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_font = Font(name="Meiryo", bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    wrap_align = Alignment(wrap_text=True, vertical="top")

    # Category fills
    cat_fills = {
        "核心": PatternFill(start_color="FDECEA", end_color="FDECEA", fill_type="solid"),
        "キャラ": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
        "世界観": PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid"),
        "ミスリード": PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid"),
    }

    # Headers
    headers = [
        "ID", "伏線名", "重要度", "分類",
        "提示（話数）", "提示内容",
        "中間ヒント（話数）", "中間ヒント内容",
        "回収（話数）", "回収内容",
        "関連伏線", "ステータス",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Sample rows (empty template with IDs)
    categories = ["核心", "キャラ", "世界観", "ミスリード"]
    for i in range(1, 13):
        row = i + 1
        cat = categories[(i - 1) % len(categories)]
        fill = cat_fills[cat]
        stars = "★★★" if cat in ("核心", "世界観") else ("★★" if cat == "キャラ" else "★")

        row_data = [
            f"F-{i:03d}", "", stars, cat,
            "", "", "", "", "", "", "", "未回収",
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.alignment = wrap_align
            cell.border = thin_border
            cell.font = Font(name="Meiryo", size=10)
            cell.fill = fill

    # Column widths
    widths = [8, 20, 8, 10, 10, 40, 12, 40, 10, 50, 15, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L13"

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    print(f"  作成: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(f"使い方: {sys.argv[0]} <output_path> [title]")
        sys.exit(1)
    title = sys.argv[2] if len(sys.argv) > 2 else "新連載"
    generate(sys.argv[1], title)
